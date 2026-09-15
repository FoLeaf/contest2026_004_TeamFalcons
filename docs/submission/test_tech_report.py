"""报告生成回归测试，不调用固件构建、网络服务或开发板。"""

from __future__ import annotations

import copy
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageChops

import build_tech_report as build
import render_figures
import report_content as content
from render_figures import render_all


class ReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory(prefix="vg-report-tests-")
        cls.root = Path(cls.temporary.name).resolve()
        assert cls.root.parent == Path(tempfile.gettempdir()).resolve()
        cls.template = build.default_publish_dir() / build.TEMPLATE_NAME
        cls.evidence_file = build.SUB / "evidence/report-evidence.json"
        cls.evidence = json.loads(cls.evidence_file.read_text(encoding="utf-8"))
        cls.parts = content.sections(cls.evidence)
        render_all(cls.root / "figures")
        cls.docx = cls.root / "report.docx"
        cls.before = build.digest(cls.template)
        build.fill_docx(cls.template, cls.docx, cls.parts, cls.root / "figures")
        cls.doc = Document(str(cls.docx))

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def test_official_sections_and_abstract(self):
        build.validate_content(self.parts, self.evidence)
        self.assertLessEqual(len(content.ABSTRACT), 300)
        self.assertEqual(len(self.parts), 7)

    def test_bad_reference_is_rejected(self):
        parts = copy.deepcopy(self.parts)
        parts[0]["blocks"].append(content.p("缺少来源 [999]"))
        with self.assertRaisesRegex(ValueError, "Unknown reference"):
            build.validate_content(parts, self.evidence)

    def test_bad_table_shape_is_rejected(self):
        parts = copy.deepcopy(self.parts)
        next(block for block in parts[1]["blocks"] if block["kind"] == "table")["rows"][0].pop()
        with self.assertRaisesRegex(ValueError, "Invalid table row"):
            build.validate_content(parts, self.evidence)

    def test_forbidden_quotes_are_rejected(self):
        parts = copy.deepcopy(self.parts)
        parts[0]["blocks"].append(content.p("\u201cplaceholder\u201d"))
        with self.assertRaisesRegex(ValueError, "quotation"):
            build.validate_content(parts, self.evidence)

    def test_docx_has_native_tables_figures_and_links(self):
        self.assertEqual(len(self.doc.tables), 10)
        self.assertEqual(len(self.doc.inline_shapes), len(build.FIGURE_NAMES))
        self.assertGreater(len(self.doc._element.xpath(".//w:hyperlink")), 50)
        self.assertFalse(any(p.text.startswith("|") for p in self.doc.paragraphs))

    def test_complete_content_is_in_docx(self):
        text = build.normalize(build.xml_text(self.doc._element))
        for expected in build.expected_strings(self.parts):
            with self.subTest(text=expected[:40]):
                self.assertIn(build.normalize(expected), text)

    def test_template_geometry_and_identity_are_preserved(self):
        section = self.doc.sections[0]
        self.assertEqual(len(self.doc.sections), 1)
        self.assertAlmostEqual(section.left_margin.inches, 1.25)
        self.assertAlmostEqual(section.top_margin.inches, 1.0)
        self.assertEqual(section._sectPr.pgSz.get(qn("w:w")), "11905")
        self.assertEqual(section._sectPr.pgSz.get(qn("w:h")), "16840")
        self.assertEqual(self.before, build.digest(self.template))

    def test_watermark_and_template_hints_are_removed(self):
        self.assertFalse(self.doc.sections[0].header._element.xpath(".//w:drawing|.//w:pict"))
        text = build.xml_text(self.doc._element)
        for hint in ("一、需提交材料清单", "三、评审维度对照", "300 字以内", "xx%"):
            self.assertNotIn(hint, text)

    def test_untouched_template_parts_remain_byte_identical(self):
        with zipfile.ZipFile(self.template) as before, zipfile.ZipFile(self.docx) as after:
            for name in build.PRESERVE_PARTS:
                self.assertEqual(before.read(name), after.read(name), name)

    def test_diagrams_are_nonblank(self):
        # 自绘示意图统一 1800 宽；界面帧保持原始分辨率，不放大到 1800。
        drawn = dict(render_figures.DRAWN)
        for path in (self.root / "figures").glob("*.png"):
            with Image.open(path) as image:
                if path.stem in drawn:
                    self.assertEqual(image.width, 1800, path.name)
                else:
                    self.assertIn(path.stem, render_figures.STILL_FIGURES)
                    self.assertGreater(image.width, 400, path.name)
                self.assertIsNotNone(ImageChops.difference(image, Image.new("RGB", image.size, "white")).getbbox())

    def test_markdown_uses_same_content(self):
        path = self.root / "report.md"
        build.write_markdown(path, self.parts)
        text = path.read_text(encoding="utf-8")
        self.assertIn(content.ABSTRACT, text)
        self.assertIn(f"{self.evidence['source_line_inventory']['physical_lines']:,}", text)
        for part in self.parts:
            self.assertIn(part["heading"], text)
        for name in ("system-map", "agent-boundary", "data-flow"):
            self.assertIn(f"figures/{name}.png", text)

    def test_export_refuses_old_pdf(self):
        target = self.root / "old.pdf"
        target.write_bytes(b"old PDF must not count as a new export")
        with patch.object(build.subprocess, "run") as run:
            with self.assertRaises(FileExistsError):
                build.export_pdf(self.docx, target)
            run.assert_not_called()

    def test_failed_export_does_not_publish(self):
        output = self.root / "existing-output"
        output.mkdir()
        pdf = output / (build.REPORT_NAME + ".pdf")
        docx = output / (build.REPORT_NAME + ".docx")
        pdf.write_bytes(b"previous PDF")
        docx.write_bytes(b"previous DOCX")
        with patch.object(build, "export_pdf", side_effect=RuntimeError("simulated Word failure")):
            with self.assertRaisesRegex(RuntimeError, "simulated Word failure"):
                build.build(self.template, output, self.root / "published", self.evidence_file)
        self.assertEqual(pdf.read_bytes(), b"previous PDF")
        self.assertEqual(docx.read_bytes(), b"previous DOCX")
        self.assertFalse((self.root / "published").exists())
        self.assertEqual(self.before, build.digest(self.template))

    def test_changed_template_is_rejected(self):
        template = self.root / "wrong-template.docx"
        template.write_bytes(b"not the official template")
        with self.assertRaisesRegex(ValueError, "official template changed"):
            build.build(template, self.root / "out", None, self.evidence_file)

    def test_wsl_path_conversion(self):
        class LinuxPath:
            def resolve(self):
                return "/mnt/f/Documents/report.docx"
        with patch.object(build.os, "name", "posix"):
            with patch.object(build.subprocess, "check_output", return_value="F:\\Documents\\report.docx\n") as run:
                self.assertEqual(build.windows_path(LinuxPath()), "F:\\Documents\\report.docx")
                run.assert_called_once_with(
                    ["wslpath", "-w", "/mnt/f/Documents/report.docx"], text=True
                )

    def test_pdf_page_footer_does_not_split_body_comparison(self):
        first = build.pdf_body_text("段落上半部分\n第 4 页\n")
        second = build.pdf_body_text("段落下半部分\n第 5 页\n")
        self.assertEqual(build.normalize(first + second), "段落上半部分段落下半部分")

    def test_pdf_inline_page_reference_is_not_removed(self):
        self.assertEqual(build.pdf_body_text("参见第 4 页的表格"), "参见第 4 页的表格")


if __name__ == "__main__":
    unittest.main(verbosity=2)
