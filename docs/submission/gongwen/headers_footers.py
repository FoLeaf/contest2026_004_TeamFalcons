# -*- coding: utf-8 -*-
"""页眉页脚：公文式页码（单页居右 / 双页居左，宋体四号，一字线）"""
import sys
import docx
from lxml import etree

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
XML_NS = 'http://www.w3.org/XML/1998/namespace'
def w(t): return '{%s}%s' % (W_NS, t)

def mk(tag, attrs=None, kids=None, text=None):
    e = etree.Element(w(tag))
    for k, v in (attrs or {}).items():
        e.set(k if k.startswith('{') else w(k), str(v))
    for c in (kids or []): e.append(c)
    if text is not None: e.text = text
    return e

SONG = '宋体'; LATIN = 'Times New Roman'
SZ_HDR = 18   # 小五 9pt
SZ_FTR = 28   # 四号 14pt
BLACK = '000000'

def rpr(size, east=SONG, bold=False):
    r = mk('rPr')
    r.append(mk('rFonts', {'ascii': LATIN, 'hAnsi': LATIN, 'eastAsia': east, 'cs': LATIN}))
    if bold:
        r.append(mk('b')); r.append(mk('bCs'))
    r.append(mk('color', {'val': BLACK}))
    r.append(mk('sz', {'val': size})); r.append(mk('szCs', {'val': size}))
    return r

def clear(p_el):
    for c in list(p_el):
        p_el.remove(c)

def set_ppr(p_el, jc=None, border=False):
    pPr = mk('pPr')
    pPr.append(mk('widowControl'))
    if border:
        pPr.append(mk('pBdr', kids=[
            mk('bottom', {'val': 'single', 'sz': '6', 'space': '2', 'color': BLACK})]))
    pPr.append(mk('spacing', {'before': '0', 'after': '0', 'line': '240', 'lineRule': 'auto'}))
    pPr.append(mk('ind', {'left': '0', 'firstLine': '0', 'firstLineChars': '0'}))
    if jc: pPr.append(mk('jc', {'val': jc}))
    p_el.insert(0, pPr)

def page_field_runs(size):
    """返回 [begin, instr, separate, '1', end] 五个 run"""
    out = []
    r = mk('r'); r.append(rpr(size)); r.append(mk('fldChar', {'fldCharType': 'begin'})); out.append(r)
    r = mk('r'); r.append(rpr(size))
    it = mk('instrText', {'{%s}space' % XML_NS: 'preserve'}); it.text = ' PAGE '
    r.append(it); out.append(r)
    r = mk('r'); r.append(rpr(size)); r.append(mk('fldChar', {'fldCharType': 'separate'})); out.append(r)
    r = mk('r'); r.append(rpr(size)); r.append(mk('t', text='1')); out.append(r)
    r = mk('r'); r.append(rpr(size)); r.append(mk('fldChar', {'fldCharType': 'end'})); out.append(r)
    return out

def text_run(text, size, east=SONG):
    r = mk('r'); r.append(rpr(size, east)); r.append(mk('t', {'{%s}space' % XML_NS: 'preserve'}, text=text))
    return r

def build_page_footer(container, jc):
    p = container.paragraphs[0]._p
    clear(p)
    set_ppr(p, jc=jc)
    p.append(text_run('— ', SZ_FTR))
    for r in page_field_runs(SZ_FTR):
        p.append(r)
    p.append(text_run(' —', SZ_FTR))

def build_header(container):
    p = container.paragraphs[0]._p
    clear(p)
    set_ppr(p, jc='center', border=True)
    p.append(text_run('VelaGuard 技术报告', SZ_HDR, east=SONG))

def blank(container):
    p = container.paragraphs[0]._p
    clear(p)
    set_ppr(p, jc='left')

SRC = sys.argv[1]; DST = sys.argv[2] if len(sys.argv) > 2 else SRC
doc = docx.Document(SRC)
log = []

secs = doc.sections
log.append('节数: %d' % len(secs))
s = secs[0]

# 首页（标题+目录）另用一套页眉页脚：不显示页眉与页码
s.different_first_page_header_footer = True
blank(s.first_page_header)
blank(s.first_page_footer)
log.append('首页(标题+目录): 无页眉、无页码')

# 正文各页：页眉加作品名
for c in (s.header, s.even_page_header):
    c.is_linked_to_previous = False
    build_header(c)
log.append('页眉: “VelaGuard 技术报告”（宋体小五·居中·下边框）')

# 页码：正文从 1 开始（首页占 0，不显示）
s.footer.is_linked_to_previous = False
s.even_page_footer.is_linked_to_previous = False
build_page_footer(s.footer, 'right')
build_page_footer(s.even_page_footer, 'left')
log.append('页码: 宋体四号一字线；单页居右 / 双页居左')

# pgNumType start=0 → 目录页为 0（不显示），正文首页显示 1
sectPr = s._sectPr
for e in sectPr.findall(w('pgNumType')):
    sectPr.remove(e)
pgnum = mk('pgNumType', {'start': '0'})
pgSz = sectPr.find(w('pgSz'))
pgSz.addnext(pgnum)
log.append('pgNumType start=0（正文首页 = 第 1 页）')

doc.save(DST)
print('\n'.join(log)); print('saved:', DST)
